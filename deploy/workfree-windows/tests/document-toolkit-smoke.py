"""Create isolated fixtures to verify installed document/data libraries."""
import json
import sys
from pathlib import Path
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

folder = Path(sys.argv[1]).resolve()
folder.mkdir(parents=True, exist_ok=True)
digital = canvas.Canvas(str(folder / "digital.pdf"))
digital.drawString(50, 740, "Invoice total: 742 dollars")
digital.showPage()
digital.drawString(50, 740, "Second page reference: 915")
digital.save()
image = Image.new("RGB", (1600, 900), "white")
draw = ImageDraw.Draw(image)
font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 52)
draw.text((70, 130), "Invoice total: 742 dollars", fill="black", font=font)
draw.text((70, 250), "Relatorio de vendas: outubro", fill="black", font=font)
image.save(folder / "scan.png")
scan = canvas.Canvas(str(folder / "scanned.pdf"), pagesize=(800, 450))
scan.drawImage(ImageReader(image), 0, 0, width=800, height=450)
scan.save()
doc = Document()
doc.add_paragraph("Word fixture reference 812")
doc.save(folder / "document.docx")
assert "812" in Document(folder / "document.docx").paragraphs[0].text
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[6])
slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Presentation fixture reference 729"
slides.save(folder / "presentation.pptx")
assert len(Presentation(folder / "presentation.pptx").slides) == 1
workbook = Workbook()
workbook.active.append(["Revenue", 452])
workbook.save(folder / "spreadsheet.xlsx")
assert load_workbook(folder / "spreadsheet.xlsx").active["B1"].value == 452
data = pd.DataFrame({"month": ["September", "October"], "revenue": [381, 452]})
data.to_csv(folder / "analysis.csv", index=False)
assert int(data.revenue.sum()) == 833
plt.bar(data.month, data.revenue)
plt.ylabel("Revenue")
plt.savefig(folder / "chart.png")
plt.close()
print(json.dumps({"ok": True, "folder": str(folder), "files": sorted(p.name for p in folder.iterdir())}))
